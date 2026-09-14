import sys
import unittest
from pathlib import Path
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "xlsx_resolver"))
from formula_values import FormulaValueError, _scan_formulas, _scan_defined_names


class DependencySafetyTests(unittest.TestCase):
    def test_same_coordinate_on_another_sheet_is_not_self_reference(self):
        book = Workbook()
        book.active['A1'] = '=other!A1'
        book.create_sheet('Other')['A1'] = 3
        self.assertEqual(len(_scan_formulas(book)), 1)

    def test_cycles_in_cells_ranges_columns_rows_and_names_are_rejected(self):
        for mode in ('absolute', 'wide', 'column', 'row', 'name', 'local_name', 'cross_sheet'):
            with self.subTest(mode=mode):
                book = Workbook()
                ws = book.active
                if mode == 'absolute':
                    ws['A1'], ws['B1'] = '=$B$1', '=$A$1'
                elif mode == 'wide':
                    ws['AA1'], ws['AB1'] = '=SUM(AB1:AZ10000)', '=AA1'
                elif mode == 'column':
                    ws['AA1'] = '=SUM(AA:AA)'
                elif mode == 'row':
                    ws['AA1'] = '=SUM(1:1)'
                elif mode in ('name', 'local_name'):
                    owner = book if mode == 'name' else ws
                    owner.defined_names.add(DefinedName('Cycle', attr_text="'Sheet'!$A$1"))
                    ws['A1'] = '=Cycle'
                else:
                    ws['A1'] = '=Other!AA1'
                    book.create_sheet('Other')['AA1'] = '=Sheet!A1'
                with self.assertRaises(FormulaValueError):
                    _scan_formulas(book)

    def test_large_range_does_not_include_unrelated_formula(self):
        book = Workbook()
        book.active['A1'] = '=SUM(B1:C100000)'
        self.assertEqual(len(_scan_formulas(book)), 1)

    def test_unresolved_structured_reference_fails_closed(self):
        book = Workbook()
        book.active['A1'] = '=SUM(Table1[Value])'
        with self.assertRaises(FormulaValueError):
            _scan_formulas(book)

    def test_local_name_cannot_hide_volatile_or_udf(self):
        for expression in ('NOW()', 'SecretUdf(1)'):
            book = Workbook()
            book.active.defined_names.add(DefinedName('Unsafe', attr_text=expression))
            with self.assertRaises(FormulaValueError):
                _scan_defined_names(book)

    def test_long_acyclic_chain_avoids_recursion_limit(self):
        book = Workbook()
        ws = book.active
        ws['A1'] = 1
        for row in range(2, 1200):
            ws.cell(row, 1, f'=A{row-1}+1')
        self.assertEqual(len(_scan_formulas(book)), 1198)
