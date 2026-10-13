import json
import subprocess
import sys
import tempfile
import threading
import time
import unittest
import urllib.request
from pathlib import Path
from unittest.mock import patch, Mock
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).parents[1] / 'xlsx_resolver'))
import xlsx_conflict as xc
import merge_server as ms
import formula_values as fv
from prepare_progress import checkpoint, preparation_context, PreparationCancelled


class WholeSourceRepairTests(unittest.TestCase):
    def test_both_invalid_sources_repaired_in_decisions_and_whole_source_preserved(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            for side, value in [('base', 3), ('ours', 11), ('theirs', 22)]:
                book = Workbook(); ws = book.active; ws.title = 'Data'
                ws.append(['##var', 'id', 'result', 'label'])
                ws.append(['##type', 'int', 'int', 'string'])
                ws.append(['', 1, '#REF!', side]); ws.append(['', 1, value, side])
                book.save(root / f'{side}.xlsx'); book.close()
            diff = xc.diff_workbooks(root/'base.xlsx', root/'ours.xlsx', root/'theirs.xlsx', analysis_dir=root/'analysis')
            sheet = diff['sheets'][0]
            self.assertEqual(sheet['action'], 'conflict_calculation_errors')
            self.assertEqual(next(e for e in sheet['calculationErrors'] if e['side']=='ours')['defaultValue'], 0)
            entry = dict(id='Data.xlsx', path='Data.xlsx', output=str(root/'out.xlsx'), diff=diff,
                         valueSnapshots=diff['valueSnapshots'], **{side:str(root/f'{side}.xlsx') for side in ('base','ours','theirs')})
            digests = {side:xc._sha256_path(meta['valuePath']) for side,meta in entry['valueSnapshots'].items()}
            for side, expected in [('ours',11),('theirs',22)]:
                selection = {'sheet':side,'repairs':{side:{'C3':0}}}
                decisions = {'files':{'Data.xlsx':{'sheets':{'Data':selection}}}}
                xc._validate_decisions({'files':[entry]}, decisions)
                output = xc._build_merged_file(entry, decisions)
                book=load_workbook(output); self.assertEqual(book['Data']['C3'].value,0)
                self.assertEqual(book['Data']['C4'].value,expected);self.assertEqual(book['Data']['D3'].value,side)
                book.close();output.unlink()
                message=xc._build_commit_message({'files':[entry]},decisions)
                self.assertIn('resolve(sheet)',message); self.assertIn('[单元格:C3]-[值:0]',message)
                for repairs in ({}, {'C3':'0'}, {'C3':False}, {'C3':2**32}, {'C3':0,'C4':1}):
                    selection['repairs']={side:repairs}
                    with self.assertRaises(xc.ConflictError):xc._validate_decisions({'files':[entry]},decisions)
            self.assertEqual(digests,{side:xc._sha256_path(meta['valuePath']) for side,meta in entry['valueSnapshots'].items()})

    def test_declared_types_and_reference_defaults(self):
        book=Workbook();ws=book.active
        ws.append(['##var','id','ref_id','value']);ws.append(['##type','int','int','bool'])
        self.assertNotIn('defaultValue',xc._repair_field_info(ws,'B3'))
        self.assertNotIn('defaultValue',xc._repair_field_info(ws,'C3'))
        self.assertEqual(xc._repair_field_info(ws,'D3')['defaultValue'],False)
        for kind,value in [('bool',False),('string',''),('string','=literal'),('uint',0),('int',-1)]:xc._validate_repair_value(kind,value)
        for kind,value in [('bool',0),('uint',-1),('float',float('nan')),('Custom',0)]:
            with self.assertRaises(xc.ConflictError):xc._validate_repair_value(kind,value)
        book.close()


class LoadingLifecycleTests(unittest.TestCase):
    def exercise(self, mode):
        with tempfile.TemporaryDirectory() as temp:
            path=Path(temp)/'manifest.json'; gate=threading.Event();bound=threading.Event();holder={}
            real=ms.make_server
            def make(*args,**kwargs):
                server=real(*args,**kwargs);holder['server']=server;holder['url']=f'http://127.0.0.1:{server.server_port}';bound.set();return server
            def prepare(*args,**kwargs):
                checkpoint(file='fixture.xlsx',fileDone=0,fileTotal=2,phase='计算公式并验证',sheet='Data',sheetDone=0,sheetTotal=2)
                while not gate.wait(.02):checkpoint()
                if mode=='failed':raise ValueError('fixture failure')
                path.write_text(json.dumps({'version':4,'files':[{'id':'fixture','path':'fixture.xlsx','valueSnapshots':{},'diff':{'sheets':[]}}],'summary':{'fileCount':1}}))
                return path
            with patch.object(ms,'make_server',side_effect=make),patch.object(xc,'prepare_conflicts',side_effect=prepare):
                thread=threading.Thread(target=lambda:ms.serve(path,prepare_repo=temp,open_browser=False));thread.start();self.assertTrue(bound.wait(5))
                def get(route):
                    with urllib.request.urlopen(holder['url']+route,timeout=5) as r:return r.read()
                try:
                    status=json.loads(get('/api/status'));self.assertEqual(status['state'],'preparing')
                    self.assertIn(b'poll()',get('/'));self.assertFalse(path.exists())
                    if mode!='cancel':
                        gate.set()
                        for _ in range(100):
                            status=json.loads(get('/api/status'))
                            if status['state']!='preparing':break
                            time.sleep(.02)
                        self.assertEqual(status['state'],'ready' if mode=='ready' else 'failed')
                        if mode=='ready':self.assertEqual(json.loads(get('/api/diff'))['summary']['fileCount'],1)
                        else:self.assertIn('fixture failure',status['error'])
                    request=urllib.request.Request(holder['url']+'/api/shutdown',data=b'',method='POST')
                    with urllib.request.urlopen(request,timeout=5) as r:self.assertTrue(json.load(r)['success'])
                    thread.join(5);self.assertFalse(thread.is_alive())
                    if mode=='cancel':self.assertFalse(path.exists())
                finally:
                    gate.set()
                    if thread.is_alive():
                        holder['server'].shutdown();thread.join(5)
    def test_ready_transition(self):self.exercise('ready')
    def test_cancel_stops_worker_without_manifest_or_worktree_write(self):self.exercise('cancel')
    def test_failed_preparation_remains_visible_and_can_close(self):self.exercise('failed')

    def test_excel_cancel_stops_owned_harness_and_calls_owned_cleanup(self):
        cancel=threading.Event(); process=Mock()
        def communicate(timeout=None):
            if not cancel.is_set():
                cancel.set();raise subprocess.TimeoutExpired('fixture',.25)
            return '', ''
        process.communicate.side_effect=communicate
        with preparation_context(lambda p:None,cancel),patch.object(fv,'_powershell_executable',return_value='fixture'),patch.object(fv.subprocess,'Popen',return_value=process),patch.object(fv,'_cleanup_excel_process') as cleanup:
            with self.assertRaises(PreparationCancelled):fv._run_excel_recalculate(Path('source.xlsx'),Path('result.xlsx'),10)
        process.terminate.assert_called_once();cleanup.assert_called()


if __name__=='__main__':unittest.main()
