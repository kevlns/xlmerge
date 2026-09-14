"""Opt-in, real Excel verification of formula-aware three-way merges.

Run with XLSX_EXCEL_E2E=1 using the bundled Python and unittest discover.
Every workbook and Git operation is confined to TemporaryDirectory.
"""
import os
import subprocess
import sys
import tempfile
import unittest
import zipfile
from xml.etree import ElementTree as ET
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.datetime import CALENDAR_MAC_1904

RESOLVER_DIR = Path(__file__).resolve().parents[1] / "xlsx_resolver"
sys.path.insert(0, str(RESOLVER_DIR))

from xlsx_conflict import _build_merged_file, apply_manifest, load_manifest, prepare_conflicts
from formula_values import FormulaValueError, materialize_values


def git(repo, *args, check=True):
    result = subprocess.run(
        ["git", "-C", str(repo), *args], capture_output=True,
        text=True, encoding="utf-8", errors="replace", timeout=30,
    )
    if check and result.returncode:
        raise AssertionError(result.stdout + result.stderr)
    return result


def write_formula_fixture(path, side):
    book = Workbook()
    book.epoch = CALENDAR_MAC_1904
    sheet = book.active
    sheet.title = "Data"
    sheet.append(["##var", "id", "name", "major", "minor", "result", "text", "flag", "date"])
    sheet.append(["##type", "int", "string", "int", "int", "int", "string", "bool", "datetime"])
    sheet.append(["##group", "c", "c", "c", "c", "c", "c", "c", "c"])
    sheet.append(["##", "ID", "Name", "Major", "Minor", "Computed", "Text", "Flag", "Date"])
    sheet.cell(1, 10, 'empty_text')
    sheet.cell(1, 11, 'error_text')
    sheet.cell(2, 10, 'string')
    sheet.cell(2, 11, 'string')
    order = {"base": [1, 2, 3, 4, 5, 6], "ours": [1, 3, 90, 4, 5, 6], "theirs": [1, 91, 2, 3, 5, 6]}[side]
    for major in order:
        row = sheet.max_row + 1
        name = f"item-{major}"
        if (side, major) in (("ours", 4), ("theirs", 2)):
            name = f"{side}-modified-{major}"
        sheet.append([None, f"=D{row}*100+E{row}", name, major, 7,
                      f"=Parameters!$B$1*D{row}", '= "=literal"', '=1=1', '=DATE(2020,1,2)'])
        sheet.cell(row, 6).fill = PatternFill("solid", fgColor="FFABCDEF")
        sheet.cell(row, 9).number_format = "yyyy-mm-dd"
        sheet.cell(row, 10, '= ""')
        sheet.cell(row, 11, '= "#N/A"')
    parameters = book.create_sheet("Parameters")
    parameters.append(["multiplier", {"base": 10, "ours": 11, "theirs": 12}[side]])
    untouched = book.create_sheet("Untouched")
    untouched["A1"] = "=6*7"
    untouched["A2"] = "=DATE(2020,1,2)"
    untouched["A2"].number_format = "yyyy-mm-dd"
    untouched["A3"] = '= ""'
    untouched["A4"] = '= "#N/A"'
    book.save(path)
    book.close()
    # Deliberately stale source caches: affected cells must be recalculated,
    # while untouched Sheet caches must remain byte-semantically unchanged.
    with zipfile.ZipFile(path) as archive:
        parts = {entry.filename: archive.read(entry.filename) for entry in archive.infolist()}
    ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    tree = ET.fromstring(parts['xl/worksheets/sheet3.xml'])
    caches = {'A1': ('n', '999'), 'A2': ('n', '42370'), 'A3': ('str', ''), 'A4': ('str', '#N/A')}
    for cell in tree.iter(ns + 'c'):
        if cell.get('r') not in caches:
            continue
        kind, value = caches[cell.get('r')]
        cell.set('t', kind)
        cached = cell.find(ns + 'v')
        if cached is None:
            cached = ET.SubElement(cell, ns + 'v')
        cached.text = value
    parts['xl/worksheets/sheet3.xml'] = ET.tostring(tree)
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def choices_for(manifest, side):
    files = {}
    for entry in manifest["files"]:
        sheets = {}
        for sheet in entry["diff"]["sheets"]:
            choice = {
                "cells": {cell.get("id") or cell["coordinate"]: side for cell in sheet["cells"]},
                "rows": {row["id"]: "keep" for row in sheet.get("allRows", []) if row["state"] == "delete_modify"},
                "columns": {col["id"]: "keep" for col in sheet.get("columns", []) if col["state"] == "delete_modify"},
            }
            if sheet["action"].startswith("conflict_"):
                choice["sheet"] = side
            if sheet.get("renameStatus") == "rename_conflict":
                choice["rename"] = side
            if (sheet.get("mergedCells") or {}).get("status") == "conflict":
                choice["mergedCells"] = side
            sheets[sheet["name"]] = choice
        files[entry["id"]] = {"sheets": sheets}
    return {"files": files}


@unittest.skipUnless(os.environ.get("XLSX_EXCEL_E2E") == "1", "Set XLSX_EXCEL_E2E=1 for real Excel calculation")
class RealExcelMergeTests(unittest.TestCase):
    def test_calculation_errors_and_circular_references_do_not_write(self):
        with tempfile.TemporaryDirectory(prefix="xlsx-excel-error-") as directory:
            source = Path(directory) / 'source.xlsx'
            output = Path(directory) / 'output.xlsx'
            for formula in ('=1/0', '=A1+1'):
                with self.subTest(formula=formula):
                    book = Workbook()
                    book.active['A1'] = formula
                    book.save(source)
                    book.close()
                    before = source.read_bytes()
                    with self.assertRaises(FormulaValueError):
                        materialize_values(source, output)
                    self.assertEqual(source.read_bytes(), before)
                    self.assertFalse(output.exists())

    def test_formula_ids_cross_sheet_values_and_literal_writeback(self):
        with tempfile.TemporaryDirectory(prefix="xlsx-excel-e2e-") as directory:
            root = Path(directory)
            repo = root / "repo"
            repo.mkdir()
            git(repo, "init", "-b", "main")
            git(repo, "config", "user.name", "Formula Test")
            git(repo, "config", "user.email", "formula-test@example.invalid")
            path = repo / "Formula.xlsx"
            write_formula_fixture(path, "base")
            git(repo, "add", "Formula.xlsx")
            git(repo, "commit", "-m", "base")
            git(repo, "checkout", "-b", "incoming")
            write_formula_fixture(path, "theirs")
            git(repo, "commit", "-am", "remote")
            git(repo, "checkout", "main")
            write_formula_fixture(path, "ours")
            git(repo, "commit", "-am", "local")
            self.assertNotEqual(git(repo, "merge", "incoming", check=False).returncode, 0)
            before = path.read_bytes()
            index_before = git(repo, "ls-files", "--stage").stdout
            manifest_path = prepare_conflicts(repo, runtime_dir=root / "runtime")
            self.assertEqual(path.read_bytes(), before)
            self.assertEqual(git(repo, "ls-files", "--stage").stdout, index_before)
            manifest = load_manifest(manifest_path)
            entry = manifest["files"][0]
            sheet = next(s for s in entry["diff"]["sheets"] if s["name"] == "Data")
            # Same logical item moves to different physical rows in each side.
            row = next(r for r in sheet["allRows"] if r.get("baseRow") == 7)
            self.assertEqual((row["oursRow"], row["theirsRow"]), (6, 8))
            id_cell = next(c for c in sheet["allCells"] if c["rowId"] == row["id"] and c["colId"] == "id")
            self.assertEqual((id_cell["base"], id_cell["ours"], id_cell["theirs"]), (307, 307, 307))
            self.assertEqual(id_cell["resolution"], "same")
            result = next(c for c in sheet["cells"] if c["rowId"] == row["id"] and c["colId"] == "result")
            self.assertEqual((result["base"], result["ours"], result["theirs"]), (30, 33, 36))
            self.assertEqual(result["resolution"], "conflict")
            for side, multiplier in (("ours", 11), ("theirs", 12)):
                with self.subTest(side=side):
                    merged = _build_merged_file(entry, choices_for(manifest, side))
                    try:
                        output = load_workbook(merged, data_only=False)
                        try:
                            data = output["Data"]
                            values = {data.cell(r, 2).value: r for r in range(5, data.max_row + 1)}
                            self.assertEqual(set(values), {107, 207, 307, 407, 507, 607, 9007, 9107})
                            final_row = values[307]
                            self.assertEqual(data.cell(final_row, 6).value, 3 * multiplier)
                            self.assertEqual(data.cell(final_row, 6).fill.fgColor.rgb, "FFABCDEF")
                            self.assertEqual(data.cell(final_row, 7).value, "=literal")
                            self.assertEqual(data.cell(final_row, 7).data_type, "s")
                            self.assertIs(data.cell(final_row, 8).value, True)
                            self.assertEqual(data.cell(final_row, 9).value.isoformat(), "2020-01-02T00:00:00")
                            self.assertEqual(data.cell(final_row, 10).value, '')
                            self.assertEqual(data.cell(final_row, 10).data_type, 's')
                            self.assertEqual(data.cell(final_row, 11).value, '#N/A')
                            self.assertEqual(data.cell(final_row, 11).data_type, 's')
                            self.assertFalse(any(c.data_type == "f" for cells in data for c in cells))
                            self.assertEqual(output["Untouched"]["A1"].value, "=6*7")
                            self.assertEqual(output["Untouched"]["A1"].data_type, "f")
                        finally:
                            output.close()
                        cached = load_workbook(merged, data_only=True)
                        try:
                            self.assertEqual(cached['Untouched']['A1'].value, 999)
                            self.assertEqual(cached['Untouched']['A2'].value.isoformat(), '2020-01-02T00:00:00')
                            self.assertEqual(cached['Untouched']['A4'].value, '#N/A')
                            self.assertEqual(cached['Untouched']['A4'].data_type, 's')
                            self.assertEqual(cached['Untouched']['A3'].data_type, 'str')
                        finally:
                            cached.close()
                    finally:
                        merged.unlink(missing_ok=True)
            result = apply_manifest(manifest_path, choices_for(manifest, "theirs"), commit=True)
            self.assertTrue(result["success"], result)
            self.assertEqual(git(repo, "diff", "--name-only", "--diff-filter=U").stdout.strip(), "")


if __name__ == "__main__":
    unittest.main()
