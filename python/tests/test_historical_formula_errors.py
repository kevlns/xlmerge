import os
import sys
import tempfile
import unittest
import subprocess
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'xlsx_resolver'))
import formula_values as fv
import xlsx_conflict as xc


def fixture(path, side, *, insert=False):
    book = Workbook()
    ws = book.active
    ws.title = 'Data'
    ws.append(['##var', 'id', 'result', 'dependent'])
    ws.append(['##type', 'int', 'int', 'int'])
    ws.append(['', 10, '=#REF!+1' if side == 'base' else '=2+3', '=C3+1'])
    ws.append(['', 20, 7, 8])
    if insert:
        ws.append(['', 30, 9, 10])
    book.save(path)
    book.close()


def calculate(source, result, timeout):
    book = load_workbook(source)
    ws = book['Data']
    broken = '#REF!' in ws['C3'].value
    ws['C3'] = '#REF!' if broken else 5
    ws['D3'] = '#REF!' if broken else 6
    book.save(result)
    book.close()


class HistoricalFormulaErrorTests(unittest.TestCase):
    def prepare(self, root, *, insert=False):
        for side in ('base', 'ours', 'theirs'):
            fixture(root / f'{side}.xlsx', side, insert=insert and side == 'theirs')
        diff = xc.diff_workbooks(root / 'base.xlsx', root / 'ours.xlsx', root / 'theirs.xlsx', analysis_dir=root / 'analysis')
        entry = dict(id='Test.xlsx', path='Test.xlsx', output=str(root / 'output.xlsx'), diff=diff,
                     valueSnapshots=diff['valueSnapshots'],
                     **{side: str(root / f'{side}.xlsx') for side in ('base', 'ours', 'theirs')})
        return entry, {'files': [entry]}

    @patch('formula_values._run_excel_recalculate', side_effect=calculate)
    def test_unknown_and_dependent_require_choice_even_when_sides_agree(self, recalc):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            entry, manifest = self.prepare(root)
            sheet = entry['diff']['sheets'][0]
            self.assertEqual(sheet['action'], 'merge')
            cells = [cell for cell in sheet['cells'] if cell['requiresChoice']]
            self.assertEqual(len(cells), 2)
            self.assertTrue(all(cell['resolution'] == 'conflict' for cell in cells))
            self.assertEqual({error['coordinate'] for error in sheet['calculationErrors']}, {'C3', 'D3'})
            self.assertEqual(cells[0]['calculationErrors']['base']['formula'], '=#REF!+1')
            with self.assertRaises(xc.ConflictError):
                xc._validate_decisions(manifest, {'files': {}})
            for side in ('ours', 'theirs'):
                decisions = {'files': {'Test.xlsx': {'sheets': {'Data': {'cells': {cell['id']: side for cell in cells}}}}}}
                xc._validate_decisions(manifest, decisions)
                merged = xc._build_merged_file(entry, decisions)
                book = load_workbook(merged)
                self.assertEqual((book['Data']['C3'].value, book['Data']['D3'].value), (5, 6))
                self.assertTrue(all(c.data_type not in ('f', 'e') for ws in book for c in ws._cells.values()))
                book.close()
                merged.unlink()
            snapshot = load_workbook(entry['valueSnapshots']['base']['valuePath'])
            self.assertEqual(snapshot['Data']['C3'].data_type, 'e')
            snapshot.close()
            self.assertFalse((root / 'output.xlsx').exists())

    @patch('formula_values._run_excel_recalculate', side_effect=calculate)
    def test_error_with_structural_changes_requires_whole_sheet(self, recalc):
        with tempfile.TemporaryDirectory() as directory:
            entry, manifest = self.prepare(Path(directory), insert=True)
            sheet = entry['diff']['sheets'][0]
            self.assertEqual(sheet['action'], 'conflict_calculation_errors')
            self.assertTrue(sheet['requiresChoice'])
            self.assertEqual(sheet['cells'], [])
            with self.assertRaises(xc.ConflictError):
                xc._validate_decisions(manifest, {'files': {}})
            decisions = {'files': {'Test.xlsx': {'sheets': {'Data': {'sheet': 'theirs'}}}}}
            xc._validate_decisions(manifest, decisions)
            merged = xc._build_merged_file(entry, decisions)
            book = load_workbook(merged)
            self.assertEqual(book['Data'].max_row, 5)
            self.assertEqual(book['Data']['C3'].value, 5)
            book.close()
            merged.unlink()

    def test_historical_mode_keeps_unsafe_formulas_and_cycles_blocked(self):
        for formula in ('=NOW()', '=CustomUdf(1)', '=[external.xlsx]Sheet!A1', '=A1+1'):
            with self.subTest(formula=formula):
                book = Workbook()
                book.active['A1'] = formula
                with self.assertRaises(fv.FormulaValueError):
                    fv._scan_formulas(book, allow_errors=True)
                book.close()

    @patch('formula_values._run_excel_recalculate', side_effect=calculate)
    def test_whole_sheet_choice_with_rename_does_not_leave_old_sheet(self, recalc):
        with tempfile.TemporaryDirectory() as directory:
            entry, manifest = self.prepare(Path(directory), insert=True)
            path = Path(entry['valueSnapshots']['theirs']['valuePath'])
            book = load_workbook(path)
            book['Data'].title = 'Renamed'
            book.save(path)
            book.close()
            entry['valueSnapshots']['theirs']['valueSha256'] = xc._sha256_path(path)
            entry['diff']['sheets'][0]['names']['theirs'] = 'Renamed'
            decisions = {'files': {'Test.xlsx': {'sheets': {'Data': {'sheet': 'theirs'}}}}}
            merged = xc._build_merged_file(entry, decisions)
            book = load_workbook(merged)
            self.assertEqual(book.sheetnames, ['Renamed'])
            self.assertEqual(book['Renamed']['C3'].value, 5)
            book.close()
            merged.unlink()

    @patch('formula_values._run_excel_recalculate', side_effect=calculate)
    def test_current_side_error_can_be_reviewed_but_not_selected(self, recalc):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for side in ('base', 'ours', 'theirs'):
                fixture(root / f'{side}.xlsx', 'base' if side == 'ours' else 'ours')
            diff = xc.diff_workbooks(root / 'base.xlsx', root / 'ours.xlsx', root / 'theirs.xlsx', analysis_dir=root / 'analysis')
            sheet = diff['sheets'][0]
            cells = [cell for cell in sheet['cells'] if cell['requiresChoice']]
            self.assertEqual(len(cells), 2)
            self.assertTrue(all(set(cell['errorSides']) == {'ours'} for cell in cells))
            entry = dict(id='Test.xlsx', path='Test.xlsx', output=str(root / 'output.xlsx'), diff=diff,
                         valueSnapshots=diff['valueSnapshots'],
                         **{side: str(root / f'{side}.xlsx') for side in ('base', 'ours', 'theirs')})
            for side in ('ours', 'theirs'):
                decisions = {'files': {'Test.xlsx': {'sheets': {'Data': {'cells': {cell['id']: side for cell in cells}}}}}}
                if side == 'ours':
                    with self.assertRaisesRegex(xc.ConflictError, 'Selected ours'):
                        xc._validate_decisions({'files': [entry]}, decisions)
                    with self.assertRaisesRegex(xc.ConflictError, 'Excel error'):
                        xc._build_merged_file(entry, decisions)
                else:
                    xc._validate_decisions({'files': [entry]}, decisions)
                    merged = xc._build_merged_file(entry, decisions)
                    book = load_workbook(merged)
                    self.assertEqual(book['Data']['C3'].value, 5)
                    book.close()
                    merged.unlink()
            self.assertFalse((root / 'output.xlsx').exists())

    @patch('formula_values._run_excel_recalculate', side_effect=calculate)
    def test_invalid_record_identity_requires_whole_sheet(self, recalc):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for side in ('base', 'ours', 'theirs'):
                path = root / f'{side}.xlsx'
                fixture(path, side)
                if side == 'base':
                    book = load_workbook(path)
                    book['Data']['B3'] = '#REF!'
                    book.save(path)
                    book.close()
            diff = xc.diff_workbooks(root / 'base.xlsx', root / 'ours.xlsx', root / 'theirs.xlsx', analysis_dir=root / 'analysis')
            self.assertEqual(diff['sheets'][0]['action'], 'conflict_calculation_errors')

    @patch('formula_values._run_excel_recalculate', side_effect=calculate)
    def test_both_sides_errors_use_declared_defaults_and_preserve_sources(self, recalc):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for side in ('base', 'ours', 'theirs'):
                fixture(root / f'{side}.xlsx', 'base')
            original = {side: (root / f'{side}.xlsx').read_bytes() for side in ('base', 'ours', 'theirs')}
            diff = xc.diff_workbooks(root / 'base.xlsx', root / 'ours.xlsx', root / 'theirs.xlsx', analysis_dir=root / 'analysis')
            sheet = diff['sheets'][0]
            self.assertEqual(sheet['action'], 'merge')
            self.assertTrue(sheet['rebuildRequired'])
            cells = [cell for cell in sheet['cells'] if cell.get('replacement')]
            self.assertEqual(len(cells), 2)
            self.assertTrue(all(not cell['requiresChoice'] and not cell['selectable'] for cell in cells))
            entry = dict(id='Test.xlsx', path='Test.xlsx', output=str(root / 'output.xlsx'), diff=diff,
                         valueSnapshots=diff['valueSnapshots'],
                         **{side: str(root / f'{side}.xlsx') for side in ('base', 'ours', 'theirs')})
            decisions = {'files': {}}
            xc._validate_decisions({'files': [entry]}, decisions)
            merged = xc._build_merged_file(entry, decisions)
            book = load_workbook(merged)
            self.assertEqual((book['Data']['C3'].value, book['Data']['D3'].value), (0, 0))
            book.close()
            merged.unlink()
            for side, content in original.items():
                self.assertEqual((root / f'{side}.xlsx').read_bytes(), content)

    def test_defaults_require_matching_types_unique_ids_and_no_references(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for side in ('base', 'ours', 'theirs'):
                book = Workbook()
                ws = book.active
                ws.title = 'Data'
                ws.append(['##var', 'id', 'count', 'text', 'enabled', 'heroId', 'custom', 'unknown', 'target'])
                ws.append(['##type', 'int', 'int', 'string', 'bool', 'int', 'MyType', None, 'int'])
                ws.append(['##ref', None, None, None, None, None, None, None, 'TargetTable'])
                ws.append(['', 10, *(['=#REF!'] * 7)])
                book.save(root / f'{side}.xlsx')
                book.close()
            def error_results(source, result, timeout):
                book = load_workbook(source)
                for ws in book:
                    for cell in ws._cells.values():
                        if cell.data_type == 'f':
                            cell.value = '#REF!'
                book.save(result)
                book.close()
            with patch('formula_values._run_excel_recalculate', side_effect=error_results):
                diff = xc.diff_workbooks(root / 'base.xlsx', root / 'ours.xlsx', root / 'theirs.xlsx', analysis_dir=root / 'analysis')
            for side in ('ours', 'theirs'):
                meta = diff['valueSnapshots'][side]
                self.assertEqual({repair['field'] for repair in meta['replacements']}, {'count', 'text', 'enabled'})
                book = load_workbook(meta['valuePath'])
                self.assertEqual((book['Data']['C4'].value, book['Data']['D4'].value, book['Data']['E4'].value), (0, '', False))
                self.assertEqual((book['Data']['C4'].data_type, book['Data']['D4'].data_type, book['Data']['E4'].data_type), ('n', 's', 'b'))
                book.close()
                self.assertEqual({error['coordinate'] for error in meta['errors']}, {'F4', 'G4', 'H4', 'I4'})

    def test_duplicate_ids_and_mismatched_types_do_not_get_defaults(self):
        for duplicate, kind in ((True, 'int'), (False, 'string')):
            with self.subTest(duplicate=duplicate), tempfile.TemporaryDirectory() as directory:
                root = Path(directory)
                for side in ('ours', 'theirs'):
                    fixture(root / f'{side}.xlsx', 'base')
                    book = load_workbook(root / f'{side}.xlsx')
                    if duplicate:
                        book['Data']['B4'] = 10
                    elif side == 'theirs':
                        book['Data']['C2'] = kind
                        book['Data']['D2'] = kind
                    book.save(root / f'{side}.xlsx')
                    book.close()
                with patch('formula_values._run_excel_recalculate', side_effect=calculate):
                    paths, meta = xc._snapshot_workbooks(None, root / 'ours.xlsx', root / 'theirs.xlsx', root / 'analysis')
                self.assertFalse(any(side.get('replacements') for side in meta.values()))

    def test_broken_sheet_references_have_no_dependency_but_other_cycles_still_fail(self):
        book = Workbook()
        book.active['A1'] = "='Deleted sheet'!#REF!+1"
        self.assertEqual(len(fv._scan_formulas(book, allow_errors=True)), 1)
        book.active['B1'] = '=C1+1'
        book.active['C1'] = '=B1+1'
        with self.assertRaises(fv.FormulaValueError):
            fv._scan_formulas(book, allow_errors=True)
        book.close()

    def test_sparse_formula_scan_does_not_allocate_formatted_rectangle(self):
        book = Workbook()
        ws = book.active
        ws['A1'] = '=2+3'
        ws['XFD1048576'] = 'last cell'
        self.assertEqual(len(fv._scan_formulas(book)), 1)
        self.assertEqual(len(ws._cells), 2)
        book.close()

    def test_ui_requires_explicit_choice_and_reset_restores_pending_state(self):
        html = (Path(xc.__file__).parent / 'merge_ui.html').read_text(encoding='utf-8')
        script = html.split('<script>')[1].split('</script>')[0].split("byId('allLocalBtn').onclick")[0]
        checks = """
const assert = require('node:assert/strict');
const pending = {id: 'r|c', coordinate: 'C3', resolution: 'conflict', requiresChoice: true};
const sheet = {name: 'Data', action: 'merge', cells: [pending]};
const file = {id: 'Test.xlsx', diff: {sheets: [sheet]}};
model = {files: [file]};
initializeDecisions();
assert.equal(sheetDecision(file, sheet).cells[pending.id], null);
assert.equal(defaultSheetChoice({action: 'conflict_calculation_errors', requiresChoice: true}), null);
renderTable = renderConflictNavigation = updateStats = render = () => {};
cycleCell(file, sheet, pending);
assert.equal(sheetDecision(file, sheet).cells[pending.id], 'ours');
cycleCell(file, sheet, pending);
assert.equal(sheetDecision(file, sheet).cells[pending.id], 'theirs');
resetSheet();
assert.equal(sheetDecision(file, sheet).cells[pending.id], null);
assert.equal(sheetDecision(file, sheet).confirmed, false);
"""
        result = subprocess.run(['node'], input=script + checks, text=True, encoding='utf-8', capture_output=True)
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_malformed_or_missing_results_are_not_unknown_values(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            fixture(root / 'base.xlsx', 'base')
            def missing(source, result, timeout):
                # Saving formulas without calculating leaves no valid caches.
                book = load_workbook(source)
                book.save(result)
                book.close()
            with patch('formula_values._run_excel_recalculate', side_effect=missing):
                with self.assertRaisesRegex(fv.FormulaValueError, 'cached result'):
                    fv.materialize_values(root / 'base.xlsx', root / 'out.xlsx', allow_errors=True)
            self.assertFalse((root / 'out.xlsx').exists())

    @unittest.skipUnless(os.environ.get('XLSX_EXCEL_E2E') == '1', 'Requires real Excel')
    def test_real_excel_historical_error_and_dependency(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            fixture(root / 'base.xlsx', 'base')
            before = (root / 'base.xlsx').read_bytes()
            with self.assertRaises(fv.FormulaValueError):
                fv.materialize_values(root / 'base.xlsx', root / 'strict.xlsx')
            fv.materialize_values(root / 'base.xlsx', root / 'history.xlsx', allow_errors=True)
            book = load_workbook(root / 'history.xlsx')
            self.assertEqual((book['Data']['C3'].data_type, book['Data']['D3'].data_type), ('e', 'e'))
            book.close()
            self.assertEqual((root / 'base.xlsx').read_bytes(), before)
            self.assertFalse((root / 'strict.xlsx').exists())

    @unittest.skipUnless(os.environ.get('XLSX_EXCEL_E2E') == '1', 'Requires real Excel')
    def test_real_excel_defaults_are_literal_audited_output(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for side in ('base', 'ours', 'theirs'):
                fixture(root / f'{side}.xlsx', 'base')
            diff = xc.diff_workbooks(root / 'base.xlsx', root / 'ours.xlsx', root / 'theirs.xlsx', analysis_dir=root / 'analysis')
            entry = dict(id='Test.xlsx', path='Test.xlsx', output=str(root / 'output.xlsx'), diff=diff,
                         valueSnapshots=diff['valueSnapshots'],
                         **{side: str(root / f'{side}.xlsx') for side in ('base', 'ours', 'theirs')})
            decisions = {'files': {}}
            xc._validate_decisions({'files': [entry]}, decisions)
            merged = xc._build_merged_file(entry, decisions)
            book = load_workbook(merged)
            self.assertEqual((book['Data']['C3'].value, book['Data']['D3'].value), (0, 0))
            self.assertTrue(all(c.data_type not in ('f', 'e') for ws in book for c in ws._cells.values()))
            book.close()
            merged.unlink()
            self.assertIn('resolve(default)', xc._build_commit_message({'files': [entry]}, decisions))


if __name__ == '__main__':
    unittest.main()
