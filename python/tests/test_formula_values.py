import os
import tempfile
import unittest
import zipfile
import hashlib
from xml.etree import ElementTree as ET
from datetime import datetime
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName

RESOLVER_DIR = Path(__file__).resolve().parents[1] / "xlsx_resolver"
import sys
sys.path.insert(0, str(RESOLVER_DIR))

import formula_values as fv


def shared_workbook(directory, mutations=None):
    """Small OOXML fixture with mixed references and a two-dimensional group."""
    source = Path(directory) / "shared.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A2"], ws["A3"], ws["A4"] = 2, 3, 9
    ws["C1"], ws["D1"] = 8, 4
    for coordinate in ("B2", "C2", "B3", "C3"):
        ws[coordinate] = "=1"
    wb.save(source)
    wb.close()
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    rewritten = Path(directory) / "rewritten.xlsx"
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(rewritten, "w") as dst:
        for item in src.infolist():
            content = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(content)
                nodes = {cell.get("r"): cell.find(ns + "f") for cell in root.iter(ns + "c")}
                for coordinate in ("B2", "C2", "B3", "C3"):
                    formula = nodes[coordinate]
                    formula.attrib = {"t": "shared", "si": "0"}
                    formula.text = None
                nodes["B2"].set("ref", "B2:C3")
                nodes["B2"].text = "A2+$D$1+C$1+$A3"
                if mutations:
                    mutations(nodes)
                content = ET.tostring(root)
            dst.writestr(item, content)
    rewritten.replace(source)
    return source


class FormulaValuesTests(unittest.TestCase):
    def test_no_formula_returns_source_without_writing_output(self):
        with tempfile.TemporaryDirectory() as d:
            source = Path(d) / "source.xlsx"
            output = Path(d) / "output.xlsx"
            wb = Workbook()
            wb.active["A1"] = "=literal"
            wb.active["A1"].data_type = "s"
            wb.save(source)
            wb.close()
            # openpyxl intentionally stores this as a literal string.
            self.assertEqual(fv.materialize_values(source, output), source.resolve())
            self.assertFalse(output.exists())

    def test_literals_preserve_types_and_equals_prefix(self):
        with tempfile.TemporaryDirectory() as d:
            source = Path(d) / "source.xlsx"
            calculated = Path(d) / "calculated.xlsx"
            output = Path(d) / "output.xlsx"
            wb = Workbook()
            ws = wb.active
            ws["A1"] = "=1+2"
            ws["A2"] = "=1=1"
            ws["A3"] = '= "=literal"'
            ws["A4"] = "=DATE(2020,1,2)"
            ws["A4"].number_format = "yyyy-mm-dd"
            wb.save(source)
            wb.close()
            # This stands in for Excel's recalculated cache and exercises the
            # final, safety-critical typed replacement independently of COM.
            calculated_wb = Workbook()
            calc = calculated_wb.active
            calc["A1"] = 3
            calc["A2"] = True
            calc["A3"] = "=literal"
            calc["A3"].data_type = "s"
            calc["A4"] = datetime(2020, 1, 2)
            calc["A4"].number_format = "yyyy-mm-dd"
            calculated_wb.save(calculated)
            calculated_wb.close()
            fv._materialize_from_result(source, calculated, output,
                                         {("Sheet", "A1"), ("Sheet", "A2"), ("Sheet", "A3"), ("Sheet", "A4")})
            result = load_workbook(output, data_only=False)
            ws = result.active
            self.assertEqual(ws["A1"].value, 3)
            self.assertIs(ws["A2"].value, True)
            self.assertEqual(ws["A3"].value, "=literal")
            self.assertEqual(ws["A3"].data_type, "s")
            self.assertEqual(ws["A4"].value, datetime(2020, 1, 2))
            self.assertEqual(ws["A4"].number_format, "yyyy-mm-dd")
            self.assertFalse(any(c.data_type == "f" for row in ws.iter_rows() for c in row))
            result.close()

    def test_unknown_udf_and_volatile_formula_fail_before_excel(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "=MyCompanyFunction(1)"
        with self.assertRaises(fv.FormulaValueError):
            fv._scan_formulas(wb)
        ws["A1"] = "=NOW()"
        with self.assertRaises(fv.FormulaValueError):
            fv._scan_formulas(wb)
        wb.close()

    def test_defined_name_formula_is_scanned_for_volatile_and_udf(self):
        wb = Workbook()
        wb.defined_names.add(DefinedName("UnsafeVolatile", attr_text="=NOW()"))
        with self.assertRaises(fv.FormulaValueError):
            fv._scan_defined_names(wb)
        wb.defined_names.clear()
        wb.defined_names.add(DefinedName("UnsafeUdf", attr_text="=MyUdf()"))
        with self.assertRaises(fv.FormulaValueError):
            fv._scan_defined_names(wb)
        wb.close()

    def test_formula_dependency_cycle_fails_closed(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "=B1+1"
        ws["B1"] = "=A1+1"
        with self.assertRaises(fv.FormulaValueError):
            fv._scan_formulas(wb)
        wb.close()

    def test_shared_formulas_expand_mixed_references_before_scanning(self):
        with tempfile.TemporaryDirectory() as d:
            source = shared_workbook(d)
            fv._scan_package(source)
            wb = load_workbook(source)
            try:
                formulas = {(sheet, cell): formula for sheet, cell, formula in fv._scan_formulas(wb)}
                self.assertEqual(formulas, {
                    ("Sheet", "B2"): "=A2+$D$1+C$1+$A3",
                    ("Sheet", "C2"): "=B2+$D$1+D$1+$A3",
                    ("Sheet", "B3"): "=A3+$D$1+C$1+$A4",
                    ("Sheet", "C3"): "=B3+$D$1+D$1+$A4",
                })
            finally:
                wb.close()

    def test_shared_formula_groups_fail_closed_before_excel(self):
        cases = [
            ("missing index", lambda n: n["B2"].attrib.pop("si"), "index"),
            ("missing master", lambda n: setattr(n["B2"], "text", None), "exactly one master"),
            ("duplicate master", lambda n: setattr(n["C2"], "text", "1"), "exactly one master"),
            ("outside range", lambda n: n["B2"].set("ref", "B2:B3"), "outside"),
            ("reversed range", lambda n: n["B2"].set("ref", "C3:B2"), "range"),
            ("out of bounds", lambda n: n["B2"].set("ref", "B2:XFE3"), "range"),
            ("missing range", lambda n: n["B2"].attrib.pop("ref"), "range"),
            ("follower range", lambda n: n["C2"].set("ref", "B2:C3"), "follower"),
            ("late master", lambda n: (setattr(n["B2"], "text", None),
                                      setattr(n["C2"], "text", "1")), "follows"),
        ]
        for label, mutation, error in cases:
            with self.subTest(label=label), tempfile.TemporaryDirectory() as d:
                source = shared_workbook(d, mutation)
                with mock.patch.object(fv, "_run_excel_recalculate") as excel:
                    with self.assertRaisesRegex(fv.FormulaValueError, error):
                        fv.materialize_values(source, Path(d) / "values.xlsx")
                    excel.assert_not_called()

    def test_shared_unsafe_formula_is_rejected_after_expansion(self):
        for formula in ("NOW()", "CompanyUdf(A2)", "[external.xlsx]Sheet!A1", "B2"):
            with self.subTest(formula=formula), tempfile.TemporaryDirectory() as d:
                source = shared_workbook(d, lambda n: setattr(n["B2"], "text", formula))
                with mock.patch.object(fv, "_run_excel_recalculate") as excel:
                    with self.assertRaises(fv.FormulaValueError):
                        fv.materialize_values(source, Path(d) / "values.xlsx")
                    excel.assert_not_called()

    def test_shared_snapshot_materializes_every_member_without_modifying_source(self):
        with tempfile.TemporaryDirectory() as d:
            source = shared_workbook(d)
            original = source.read_bytes()
            output = Path(d) / "values.xlsx"

            def calculate(temporary_source, result, timeout):
                self.assertNotEqual(temporary_source, source)
                wb = load_workbook(temporary_source)
                for coordinate, value in {"B2": 17, "C2": 28, "B3": 24, "C3": 41}.items():
                    wb.active[coordinate] = value
                wb.save(result)
                wb.close()

            with mock.patch.object(fv, "_run_excel_recalculate", side_effect=calculate):
                fv.materialize_values(source, output)
            self.assertEqual(source.read_bytes(), original)
            wb = load_workbook(output)
            try:
                self.assertEqual([wb.active[c].value for c in ("B2", "C2", "B3", "C3")], [17, 28, 24, 41])
                self.assertFalse(list(fv._iter_formula_cells(wb)))
            finally:
                wb.close()

    def test_array_and_data_table_formulas_remain_rejected(self):
        for kind in ("array", "dataTable"):
            with self.subTest(kind=kind), tempfile.TemporaryDirectory() as d:
                source = shared_workbook(d, lambda n: n["B2"].set("t", kind))
                with self.assertRaisesRegex(fv.FormulaValueError, "Unsupported"):
                    fv._scan_package(source)

@unittest.skipUnless(os.environ.get("XLSX_EXCEL_E2E") == "1", "set XLSX_EXCEL_E2E=1")
class ExcelIntegrationTests(unittest.TestCase):
    def test_real_excel_shared_formula_snapshot(self):
        with tempfile.TemporaryDirectory() as d:
            source = shared_workbook(d)
            digest = hashlib.sha256(source.read_bytes()).hexdigest()
            output = Path(d) / "values.xlsx"
            fv.materialize_values(source, output)
            self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), digest)
            wb = load_workbook(output)
            try:
                self.assertEqual([wb.active[c].value for c in ("B2", "C2", "B3", "C3")], [17, 28, 24, 41])
                self.assertFalse(list(fv._iter_formula_cells(wb)))
            finally:
                wb.close()

    def test_real_excel_snapshot_types(self):
        if not Path(r"C:\Program Files\Microsoft Office\Root\Office16\EXCEL.EXE").is_file():
            self.skipTest("Office16 Excel is not installed")
        with tempfile.TemporaryDirectory() as d:
            source = Path(d) / "source.xlsx"
            output = Path(d) / "snapshot.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Parameters"
            ws["B1"] = 2
            ws["D5"] = 3
            ws["E5"] = 4
            ws["F1"] = "=D5*100+E5"
            ws["F2"] = "=Parameters!$B$1*D5"
            ws["F3"] = '= "=literal"'
            ws["F4"] = "=1=1"
            ws["F5"] = "=DATE(2020,1,2)"
            ws["F5"].number_format = "yyyy-mm-dd"
            ws["F6"] = '= ""'
            wb.save(source)
            wb.close()
            self.assertEqual(fv.materialize_values(source, output), output.resolve())
            result = load_workbook(output, data_only=False)
            ws = result["Parameters"]
            self.assertEqual(ws["F1"].value, 304)
            self.assertEqual(ws["F2"].value, 6)
            self.assertEqual(ws["F3"].value, "=literal")
            self.assertIs(ws["F4"].value, True)
            self.assertEqual(ws["F5"].value, datetime(2020, 1, 2))
            self.assertEqual(ws["F6"].value, "")
            self.assertFalse(any(c.data_type == "f" for row in ws.iter_rows() for c in row))
            result.close()


if __name__ == "__main__":
    unittest.main()
